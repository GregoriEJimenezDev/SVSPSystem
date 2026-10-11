"""Tests de departamentos (CRUD + filtro) y reporte Word."""

import io

from docx import Document
from fastapi.testclient import TestClient

from app.main import app

client = TestClient(app)

ADMIN = {
    "email": "depto-admin@svsp.com",
    "full_name": "Admin Deptos",
    "password": "DeptoAdmin123!",
    "modules": ["asistencia", "adminpro"],
}


def _auth_headers(email: str = ADMIN["email"], password: str = ADMIN["password"]) -> dict:
    client.post("/api/v1/auth/register", json=ADMIN)
    r = client.post("/api/v1/auth/login", json={"email": email, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def test_department_crud():
    h = _auth_headers()
    # crear
    r = client.post("/api/v1/departments/", json={"name": "Laboratorio"}, headers=h)
    assert r.status_code == 201, r.text
    dept_id = r.json()["id"]
    # duplicado (case-insensitive) -> 400
    r = client.post("/api/v1/departments/", json={"name": "  laboratorio "}, headers=h)
    assert r.status_code == 400, r.text
    # listar contiene el creado
    names = [d["name"] for d in client.get("/api/v1/departments/", headers=h).json()]
    assert "Laboratorio" in names
    # renombrar
    r = client.put(f"/api/v1/departments/{dept_id}", json={"name": "Lab Central"}, headers=h)
    assert r.status_code == 200, r.text
    assert r.json()["name"] == "Lab Central"
    # inhabilitar (sin empleados) -> 204
    r = client.delete(f"/api/v1/departments/{dept_id}", headers=h)
    assert r.status_code == 204, r.text
    names = [d["name"] for d in client.get("/api/v1/departments/", headers=h).json()]
    assert "Lab Central" not in names


def test_department_delete_blocked_when_in_use():
    h = _auth_headers()
    dept_id = client.post("/api/v1/departments/", json={"name": "Planta"}, headers=h).json()["id"]
    client.post("/api/v1/employees/", json={"name": "Obrero Uno", "department": "Planta"}, headers=h)
    r = client.delete(f"/api/v1/departments/{dept_id}", headers=h)
    assert r.status_code == 400, r.text


def test_employees_filter_by_department_and_edit():
    h = _auth_headers()
    client.post("/api/v1/employees/", json={"name": "Ana Filtro", "department": "Norte"}, headers=h)
    client.post("/api/v1/employees/", json={"name": "Luis Filtro", "department": "Sur"}, headers=h)
    r = client.get("/api/v1/employees/?department=norte", headers=h)
    assert r.status_code == 200, r.text
    names = [e["name"] for e in r.json()]
    assert "Ana Filtro" in names
    assert "Luis Filtro" not in names
    # editar nombre + departamento
    emp_id = next(e["id"] for e in r.json() if e["name"] == "Ana Filtro")
    r = client.put(
        f"/api/v1/employees/{emp_id}",
        json={"name": "Ana Editada", "department": "Norte"},
        headers=h,
    )
    assert r.status_code == 200, r.text
    assert r.json()["name"] == "Ana Editada"


def test_report_docx_download():
    h = _auth_headers()
    client.post("/api/v1/employees/", json={"name": "Pedro Word", "department": "Oficina"}, headers=h)
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["empleado", "hora_entrada", "hora_salida"])
    ws.append(["Pedro Word", "08:15", "17:30"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    client.post(
        "/api/v1/attendance/import-excel",
        files={"file": ("a.xlsx", buf.getvalue(),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=h,
    )
    r = client.post(
        "/api/v1/attendance/report-docx",
        json={"department": "Oficina"},
        headers=h,
    )
    assert r.status_code == 200, r.text[:200]
    assert "officedocument.wordprocessingml" in r.headers["content-type"]
    doc = Document(io.BytesIO(r.content))
    table = doc.tables[0]
    headers = [c.text for c in table.rows[0].cells]
    assert headers == ["Empleado", "Hora de entrada", "Hora salida", "Total horas trabajadas"]
    body = [[c.text for c in row.cells] for row in table.rows[1:]]
    assert any("Pedro Word" in row[0] for row in body)
    assert any(row[1].startswith("08:15") for row in body if "Pedro Word" in row[0])
